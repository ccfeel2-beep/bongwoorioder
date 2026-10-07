import streamlit as st
import pandas as pd
import json
import os
import string
import re
from io import BytesIO
import io
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows

st.set_page_config(page_title="봉우리 스마트 발주 통합 시스템", layout="wide")

# --- 1. 설정 데이터 로드 및 초기화 (호환성 유지) ---
DB_FILE = "web_settings.json"

def load_settings():
    default_settings = {
        "input_mappings": {},
        "output_mapping": {
            "수취인명": "A", "수취인연락처": "B", "수취인주소": "E", 
            "수량": "F", "상품명": "D", "배송메세지": "H", 
            "주문자명": "J", "주문자연락처": "K", "취합업체": "M",
            "번호": "N", "보내는 주소": "O", "업체명2": "P"
        },
        "fixed_values": {
            "sender_address": "경기 남양주시 화도읍 답내리 220-2(디더블유대)"
        },
        "google_sheet_url": ""
    }
    
    if os.path.exists(DB_FILE):
        with open(DB_FILE, "r", encoding="utf-8") as f:
            data = json.load(f)
            for k, v in default_settings["output_mapping"].items():
                if k not in data.get("output_mapping", {}):
                    data.setdefault("output_mapping", {})[k] = v
            if "fixed_values" not in data:
                data["fixed_values"] = default_settings["fixed_values"]
            return data
    return default_settings

def save_settings(data):
    with open(DB_FILE, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=4)

if 'settings' not in st.session_state:
    st.session_state.settings = load_settings()

# --- 알파벳 열 번호 변환 함수 ---
def col_to_num(col_str):
    if not col_str: return None
    clean_str = re.sub(r'[^A-Za-z]', '', str(col_str)).upper()
    if not clean_str: return None
    
    num = 0
    for c in clean_str:
        num = num * 26 + (ord(c) - ord('A') + 1)
    return num - 1

def num_to_col(n):
    if n is None or n < 0: return ""
    string_col = ""
    n += 1
    while n > 0:
        n, remainder = divmod(n - 1, 26)
        string_col = chr(65 + remainder) + string_col
    return string_col

# --- 구글 시트 상품명 사전 불러오기 ---
@st.cache_data(ttl=60)
def load_product_dict(url):
    if not url: return {}
    try:
        df = pd.read_csv(url)
        return dict(zip(df.iloc[:, 0].astype(str).str.strip(), df.iloc[:, 1].astype(str).str.strip()))
    except Exception as e:
        st.error(f"구글 시트를 불러오는 데 실패했습니다: {e}")
        return {}

# --- 메인 화면 탭 구성 ---
st.title("🍎 봉우리 스마트 발주 통합 시스템 (Web Ver.)")

# 기존 4개 탭에서 5개 탭으로 확장
tab1, tab2, tab3, tab4, tab5 = st.tabs([
    "📦 대량 취합 실행", 
    "⚙️ 거래처 양식 관리", 
    "🛠️ 최종 출력 양식 관리", 
    "🔗 구글 시트 연동",
    "🚚 송장 정돈 및 검수"
])

# ==========================================
# 탭 3: 최종 출력 양식 관리
# ==========================================
with tab3:
    st.header("🛠️ 최종 출력 양식 관리 (도착지 열 설정)")
    st.write("최종적으로 만들어질 엑셀 파일에서 각 데이터가 어느 열(A, B, C...)에 들어갈지 설정합니다.")
    
    out_cols = st.columns(3)
    new_output = {}
    
    for i, (field, current_col) in enumerate(st.session_state.settings["output_mapping"].items()):
        with out_cols[i % 3]:
            new_output[field] = st.text_input(f"{field} (최종 열)", value=current_col, key=f"out_{field}").upper()
            
    st.write("---")
    st.subheader("📌 텍스트 고정값 설정")
    new_sender_address = st.text_input("보내는 주소에 들어갈 문구", value=st.session_state.settings["fixed_values"]["sender_address"])
            
    if st.button("💾 최종 양식 저장"):
        st.session_state.settings["output_mapping"] = new_output
        st.session_state.settings["fixed_values"]["sender_address"] = new_sender_address
        save_settings(st.session_state.settings)
        st.success("최종 출력 양식 및 고정값이 업데이트되었습니다!")

# ==========================================
# 탭 4: 구글 시트 연동 설정
# ==========================================
with tab4:
    st.header("🔗 상품명 변환용 구글 시트 연결")
    st.info("구글 시트에서 [파일] -> [공유] -> [웹에 게시]를 누르고 'CSV' 형식으로 게시한 링크를 아래에 붙여넣어 주세요.")
    
    gs_url = st.text_input("구글 시트 CSV 링크", value=st.session_state.settings.get("google_sheet_url", ""))
    if st.button("💾 링크 저장 및 테스트"):
        st.session_state.settings["google_sheet_url"] = gs_url
        save_settings(st.session_state.settings)
        dict_test = load_product_dict(gs_url)
        if dict_test:
            st.success(f"연결 성공! 총 {len(dict_test)}개의 상품명 변환 규칙을 불러왔습니다.")
            with st.expander("등록된 상품명 미리보기"):
                st.dataframe(pd.DataFrame(list(dict_test.items()), columns=["원본 상품명", "변환될 상품명"]))

# ==========================================
# 탭 2: 거래처 양식 관리
# ==========================================
with tab2:
    st.header("⚙ 거래처(입력) 양식 관리")
    col_list, col_edit = st.columns([1, 2])
    
    comp_list = sorted(list(st.session_state.settings["input_mappings"].keys()))
    
    with col_list:
        st.subheader("등록된 업체 목록")
        selected_comp = st.selectbox("수정할 업체 선택 (새로 추가하려면 '신규 추가' 선택)", ["-- 신규 추가 --"] + comp_list)
        if selected_comp != "-- 신규 추가 --" and st.button("🗑️ 업체 삭제"):
            del st.session_state.settings["input_mappings"][selected_comp]
            save_settings(st.session_state.settings)
            st.rerun()
            
        st.write("---")
        st.subheader("📥 설정 파일 백업")
        current_settings_json = json.dumps(st.session_state.settings, ensure_ascii=False, indent=4)
        st.download_button(
            label="💾 현재 설정 파일(json) 다운로드",
            data=current_settings_json,
            file_name="web_settings.json",
            mime="application/json"
        )

    with col_edit:
        st.subheader("양식 상세 설정")
        
        default_comp_name = "" if selected_comp == "-- 신규 추가 --" else selected_comp
        default_start_row = 2
        if selected_comp != "-- 신규 추가 --":
            default_start_row = st.session_state.settings["input_mappings"].get(selected_comp, {}).get("시작행", 2)

        comp_name = st.text_input("업체명", value=default_comp_name, key=f"comp_name_{selected_comp}")
        start_row = st.number_input("데이터 시작 행 (숫자)", min_value=1, value=int(default_start_row), key=f"start_row_{selected_comp}")
        
        st.write("엑셀에서 데이터가 있는 열(A, B, C...) 알파벳을 입력하세요.")
        
        exclude_fields = ["취합업체", "번호", "보내는 주소", "업체명2"]
        fields = [f for f in st.session_state.settings["output_mapping"].keys() if f not in exclude_fields]
        entries = {}
        
        edit_cols = st.columns(2)
        for i, f in enumerate(fields):
            curr_val = ""
            if selected_comp != "-- 신규 추가 --":
                comp_data = st.session_state.settings["input_mappings"].get(selected_comp, {})
                c_num = comp_data.get("cols", {}).get(f)
                curr_val = num_to_col(c_num) if c_num is not None else ""
            
            with edit_cols[i % 2]:
                entries[f] = st.text_input(f, value=curr_val, key=f"in_{f}_{selected_comp}")

        if st.button("💾 업체 양식 저장"):
            if comp_name:
                mapped_cols = {k: col_to_num(v) for k, v in entries.items()}
                st.session_state.settings["input_mappings"][comp_name] = {"cols": mapped_cols, "시작행": start_row}
                save_settings(st.session_state.settings)
                st.success(f"[{comp_name}] 양식이 저장되었습니다.")
                st.rerun()
            else:
                st.warning("업체명을 입력해주세요.")

# ==========================================
# 탭 1: 대량 취합 실행
# ==========================================
with tab1:
    st.header("📦 엑셀 파일 대량 취합")
    
    uploaded_files = st.file_uploader("취합할 엑셀 파일들을 드래그해서 올려주세요", type=["xlsx", "xls"], accept_multiple_files=True)
    
    if uploaded_files:
        st.write("---")
        comp_list = sorted(list(st.session_state.settings["input_mappings"].keys()))
        
        if 'file_order' not in st.session_state:
            st.session_state.file_order = []
        if 'last_uploaded_names' not in st.session_state:
            st.session_state.last_uploaded_names = set()

        current_names = {f.name for f in uploaded_files}
        if current_names != st.session_state.last_uploaded_names:
            st.session_state.file_order = [f for f in st.session_state.file_order if f in current_names]
            for f in uploaded_files:
                if f.name not in st.session_state.file_order:
                    st.session_state.file_order.append(f.name)
            st.session_state.last_uploaded_names = current_names
            
        file_dict = {f.name: f for f in uploaded_files}
        file_company_map = {}
        
        st.write("💡 **위아래 화살표(🔼/🔽)를 눌러 최종 엑셀에서 합쳐질 파일의 순서를 변경할 수 있습니다.**")
        
        for i, fname in enumerate(st.session_state.file_order):
            col_btn, col_name, col_sel = st.columns([1, 4, 3])
            
            with col_btn:
                b1, b2 = st.columns(2)
                with b1:
                    if st.button("🔼", key=f"up_{fname}") and i > 0:
                        st.session_state.file_order[i], st.session_state.file_order[i-1] = st.session_state.file_order[i-1], st.session_state.file_order[i]
                        st.rerun()
                with b2:
                    if st.button("🔽", key=f"down_{fname}") and i < len(st.session_state.file_order) - 1:
                        st.session_state.file_order[i], st.session_state.file_order[i+1] = st.session_state.file_order[i+1], st.session_state.file_order[i]
                        st.rerun()
            with col_name:
                st.write(f"📄 {fname}")
            with col_sel:
                file_company_map[fname] = st.selectbox(f"{fname} 업체 선택", ["-- 선택 --"] + comp_list, key=f"sel_{fname}", label_visibility="collapsed")
        
        st.write("")
        if st.button("🚀 취합 시작", type="primary", use_container_width=True):
            all_rows = []
            missing_products = set()
            
            prod_dict = load_product_dict(st.session_state.settings.get("google_sheet_url", ""))
            
            with st.spinner('파일을 병합하고 상품명을 변환하는 중...'):
                for fname in st.session_state.file_order:
                    file = file_dict[fname]
                    comp = file_company_map[fname]
                    if comp == "-- 선택 --": continue
                    
                    m = st.session_state.settings["input_mappings"][comp]
                    
                    df = None
                    file.seek(0)
                    try:
                        df = pd.read_excel(file, header=None, engine='openpyxl')
                    except Exception:
                        try:
                            file.seek(0)
                            df = pd.read_excel(file, header=None, engine='xlrd')
                        except ImportError:
                            st.error("서버에 xlrd 부품이 없습니다. Streamlit 앱을 지웠다가 다시 만들어주세요!")
                            st.stop()
                        except Exception:
                            try:
                                file.seek(0)
                                html_data = file.getvalue().decode('utf-8', errors='ignore')
                                dfs = pd.read_html(html_data)
                                df = dfs[0]
                            except Exception:
                                st.error(f"⚠️ '{fname}' 파일을 읽을 수 없습니다. 엑셀 파일이 손상되었거나 알 수 없는 형식입니다.")
                                continue
                    
                    if df is None or df.empty:
                        continue
                        
                    data_df = df.iloc[m['시작행'] - 1:]
                    
                    for _, row in data_df.iterrows():
                        name_col = m['cols'].get('수취인명')
                        if name_col is None or name_col not in row.index or pd.isna(row[name_col]) or str(row[name_col]).strip() == "": 
                            continue
                        
                        row_data = {}
                        for field in st.session_state.settings["output_mapping"].keys():
                            if field in ["취합업체", "번호", "보내는 주소", "업체명2"]:
                                continue
                                
                            c = m['cols'].get(field)
                            if c is not None and c in row.index:
                                v = row[c]
                                row_data[field] = "" if pd.isna(v) else str(v).strip()
                            else:
                                row_data[field] = ""
                                
                        row_data["취합업체"] = comp
                        
                        original_prod = row_data.get("상품명", "")
                        if original_prod:
                            if original_prod in prod_dict:
                                row_data["상품명"] = prod_dict[original_prod]
                            else:
                                missing_products.add(original_prod)
                                row_data["상품명"] = original_prod
                                row_data["_is_missing"] = True
                        
                        try:
                            q_idx = m['cols'].get('수량')
                            if q_idx is not None and q_idx in row.index:
                                cnt = int(float(row[q_idx]))
                            else:
                                cnt = 1
                        except: cnt = 1
                        
                        row_data["수량"] = 1
                        for _ in range(cnt):
                            all_rows.append(row_data.copy())
                            
            if all_rows:
                out_map = st.session_state.settings["output_mapping"]
                
                max_col_num = max([col_to_num(v) for v in out_map.values() if v])
                final_cols = [num_to_col(i) for i in range(max_col_num + 1)]
                final_df = pd.DataFrame(columns=final_cols)
                
                sender_address_text = st.session_state.settings.get("fixed_values", {}).get("sender_address", "")
                
                for idx, row_data in enumerate(all_rows):
                    excel_row_num = idx + 2
                    for field, target_col in out_map.items():
                        if target_col:
                            clean_col = re.sub(r'[^A-Za-z]', '', target_col).upper()
                            if clean_col:
                                if field == "번호":
                                    val = idx + 1
                                elif field == "보내는 주소":
                                    val = sender_address_text
                                elif field == "업체명2":
                                    val = f"=M{excel_row_num}"
                                else:
                                    val = row_data.get(field, "")
                                    
                                final_df.loc[idx, clean_col] = val

                new_headers = []
                for col_letter in final_df.columns:
                    mapped_field = None
                    for field, target_col in out_map.items():
                        if target_col and re.sub(r'[^A-Za-z]', '', target_col).upper() == col_letter:
                            mapped_field = field
                            break
                            
                    if mapped_field:
                        new_headers.append(mapped_field)
                    else:
                        new_headers.append(col_letter)
                        
                final_df.columns = new_headers
                
                output = BytesIO()
                with pd.ExcelWriter(output, engine='xlsxwriter') as writer:
                    final_df.to_excel(writer, index=False)
                    
                    workbook = writer.book
                    worksheet = writer.sheets['Sheet1']
                    orange_format = workbook.add_format({'bg_color': '#F4B084'})
                    
                    max_row = len(final_df)
                    max_col = len(final_df.columns) - 1
                    worksheet.autofilter(0, 0, max_row, max_col)
                    
                    prod_col_letter = out_map.get("상품명", "")
                    clean_prod_col = re.sub(r'[^A-Za-z]', '', prod_col_letter).upper()
                    prod_col_idx = col_to_num(clean_prod_col) if clean_prod_col else None
                    
                    if prod_col_idx is not None:
                        for idx, r_data in enumerate(all_rows):
                            if r_data.get("_is_missing"):
                                worksheet.write(idx + 1, prod_col_idx, r_data.get("상품명", ""), orange_format)
                
                excel_data = output.getvalue()
                
                st.success(f"✅ 총 {len(all_rows)}건의 데이터 취합이 완료되었습니다!")
                
                if missing_products:
                    st.error("⚠️ 아래 상품명들은 구글 시트에 등록되지 않아 변환되지 않았습니다. 시트에 추가해주세요.")
                    st.code("\n".join(missing_products))
                else:
                    st.info("✨ 모든 상품명이 완벽하게 변환되었습니다.")
                
                st.download_button(
                    label="💾 통합 엑셀 파일 다운로드",
                    data=excel_data,
                    file_name="발주취합완료.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    type="primary"
                )
            else:
                st.warning("선택된 파일에서 유효한 데이터를 찾지 못했습니다.")


# ==========================================
# 탭 5: 송장 정돈 및 검수
# ==========================================
with tab5:
    st.header("🚚 송장 원본 정돈 및 연계비용 검수")
    st.write("송장 원본 엑셀을 업로드하면 필요한 열 추출/정렬 및 연계비용 강조 파일로 자동 변환합니다.")

    uploaded_file = st.file_uploader("송장 원본 엑셀 파일 업로드", type=["xlsx", "xls"], key="tab5_invoice_file")

    if uploaded_file is not None:
        try:
            # 1. 원본 데이터 읽기 (운송장번호와 주문번호는 지수표기법 방지를 위해 문자열로, 상품코드는 숫자로 로드)
            df = pd.read_excel(uploaded_file, dtype={
                '운송장번호': str, 
                '주문번호': str
            })

            # 추출할 10개 열 지정 및 순서 정의
            target_cols = [
                '운송장번호', '수하인명', '수하인기본주소', '송하인명', 
                'A', '상품명', '상품코드', '주문번호', '쇼핑몰', '연계비용'
            ]

            available_cols = [col for col in target_cols if col in df.columns]
            df_filtered = df[available_cols].copy()

            # 2. 상품코드 숫자형 변환 및 오름차순 정렬
            if '상품코드' in df_filtered.columns:
                df_filtered['상품코드'] = pd.to_numeric(df_filtered['상품코드'], errors='coerce')
                df_filtered = df_filtered.sort_values(by='상품코드', ascending=True)

            st.success(f"총 {len(df_filtered)}건의 데이터가 성공적으로 정돈되었습니다.")

            # 웹 화면 표시 (연계비용 170 초과 시 화면에서도 노란색 강조)
            def highlight_cost(row):
                if '연계비용' in row and pd.notnull(row['연계비용']) and float(row['연계비용']) > 170:
                    return ['background-color: #FFFF99'] * len(row)
                return [''] * len(row)

            st.dataframe(df_filtered.style.apply(highlight_cost, axis=1), use_container_width=True)

            # 3. openpyxl을 활용하여 엑셀 셀 서식 적용 (상품코드는 숫자, 운송장/주문번호는 텍스트)
            output = io.BytesIO()
            wb = Workbook()
            ws = wb.active
            ws.title = "송장정리"

            # 헤더 작성
            headers = list(df_filtered.columns)
            ws.append(headers)

            # 데이터 적재
            for _, r in df_filtered.iterrows():
                row_values = []
                for col in headers:
                    val = r[col]
                    if pd.isna(val):
                        row_values.append("")
                    elif col == '상품코드':
                        try:
                            row_values.append(int(val) if float(val).is_integer() else float(val))
                        except ValueError:
                            row_values.append(val)
                    else:
                        row_values.append(str(val))
                ws.append(row_values)

            # 열별 셀 서식 지정 (운송장번호/주문번호는 텍스트, 상품코드는 숫자 형식)
            for col_idx, col_name in enumerate(headers, start=1):
                if col_name in ['운송장번호', '주문번호']:
                    for row in range(2, ws.max_row + 1):
                        ws.cell(row=row, column=col_idx).number_format = '@'
                elif col_name == '상품코드':
                    for row in range(2, ws.max_row + 1):
                        ws.cell(row=row, column=col_idx).number_format = '#,##0'

            # 1행에 자동 필터(AutoFilter) 설정
            if ws.max_row >= 1 and ws.max_column >= 1:
                from openpyxl.utils import get_column_letter
                max_col_letter = get_column_letter(ws.max_column)
                ws.auto_filter.ref = f"A1:{max_col_letter}{ws.max_row}"

            # 노란색 서식 지정
            yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

            if '연계비용' in available_cols:
                cost_col_idx = available_cols.index('연계비용') + 1

                for row in range(2, ws.max_row + 1):
                    val = ws.cell(row=row, column=cost_col_idx).value
                    if val is not None:
                        try:
                            if float(val) > 170:
                                for col in range(1, len(available_cols) + 1):
                                    ws.cell(row=row, column=col).fill = yellow_fill
                        except ValueError:
                            pass

            wb.save(output)
            output.seek(0)

            # 4. 수정본 다운로드 버튼
            st.download_button(
                label="📥 정돈된 송장 엑셀 파일 다운로드",
                data=output,
                file_name="정돈된 엑셀 파일.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )

        except Exception as e:
            st.error(f"파일 처리 중 오류가 발생했습니다: {e}")
